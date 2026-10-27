"""The daily "Long restaurant waits" sheet (9.7) — an Excel file that opens 1:1 in Google Sheets.

No Google account and no key: the system builds the file, the morning Slack message carries a download link
(/export/long-waits.xlsx, behind the dashboard password). Open it in Google Sheets: Drive → New → File upload, then
open it — or File → Import in any sheet.

Tabs: Yesterday · Last week (Mondays) · Yesterday orders · Last week orders (Mondays) · How it works.
Everything sorted by city A–Z. No customer data: order numbers, restaurants, riders and times only.
"""
from __future__ import annotations

import io
from collections import Counter, defaultdict
from datetime import date, datetime, time as dtime


# ------------------------------------------------------------------ cells and formats
class Link:
    def __init__(self, text: str, url: str):
        self.text, self.url = text, url


def _rgb(h: str) -> dict:
    h = h.lstrip("#")
    return {"red": int(h[0:2], 16) / 255, "green": int(h[2:4], 16) / 255, "blue": int(h[4:6], 16) / 255}


def fmt(bg=None, bold=None, italic=None, size=None, color=None, align=None, valign=None, wrap=None, num=None, underline=None, font=None):
    """A userEnteredFormat and its field mask (only the given parts are changed)."""
    f, tf, fields = {}, {}, []
    if bg:
        f["backgroundColor"] = _rgb(bg); fields.append("backgroundColor")
    for k, v in (("bold", bold), ("italic", italic), ("fontSize", size), ("underline", underline), ("fontFamily", font)):
        if v is not None:
            tf[k] = v; fields.append("textFormat." + k)
    if color:
        tf["foregroundColor"] = _rgb(color); fields.append("textFormat.foregroundColor")
    if tf:
        f["textFormat"] = tf
    if align:
        f["horizontalAlignment"] = align; fields.append("horizontalAlignment")
    if valign:
        f["verticalAlignment"] = valign; fields.append("verticalAlignment")
    if wrap is not None:
        f["wrapStrategy"] = "WRAP" if wrap else "OVERFLOW_CELL"; fields.append("wrapStrategy")
    if num:
        f["numberFormat"] = {"type": num[0], "pattern": num[1]}; fields.append("numberFormat")
    return f, ",".join("userEnteredFormat." + x for x in fields)


VIOLET, INK, AMBER, RED, LINK = "4B3591", "1F1A2B", "B45309", "B42318", "2B5CD6"
BAND = ["FFFFFF", "F4F1FB"]
DATE_F, TIME_F = ("DATE", "ddd dd.mm."), ("TIME", "hh:mm")
INT_F, ONE_F, PCT_F = ("NUMBER", "#,##0"), ("NUMBER", "0.0"), ("PERCENT", "0%")


class Tab:
    """One tab: values plus formatting, turned into Sheets API requests by write()."""
    def __init__(self, title, ncols, color, freeze=(0, 0), grid=True):
        self.title, self.ncols, self.color, self.freeze, self.grid = title, ncols, color, freeze, grid
        self.rows, self.styles, self.borders, self.cond, self.widths, self.heights, self.filter = [], [], [], [], {}, {}, None

    def add(self, row):
        self.rows.append(list(row) + [None] * (self.ncols - len(row)))
        return len(self.rows) - 1

    def style(self, r0, r1, c0, c1, **kw):
        self.styles.append((r0, r1, c0, c1) + fmt(**kw))

    def header(self, r, labels):
        for j, x in enumerate(labels):
            self.rows[r][j] = x
        self.style(r, r + 1, 0, len(labels), bg=VIOLET, bold=True, color="FFFFFF", align="CENTER", valign="MIDDLE", wrap=True)
        self.style(r, r + 1, 0, 1, align="LEFT")
        self.heights[r] = 40

    def wait_colors(self, r0, r1, c):
        g = (r0, r1, c, c + 1)
        self.cond += [(g, "NUMBER_GREATER_THAN_EQ", ["20"], {"backgroundColor": _rgb("F4CCCC"), "textFormat": {"bold": True, "foregroundColor": _rgb(RED)}}),
                      (g, "NUMBER_BETWEEN", ["15", "19.99"], {"backgroundColor": _rgb("FCE4D6"), "textFormat": {"bold": True, "foregroundColor": _rgb(AMBER)}}),
                      (g, "NUMBER_BETWEEN", ["10", "14.99"], {"backgroundColor": _rgb("FFF2CC")})]


# ------------------------------------------------------------------ the tabs
def _ck(s: str) -> str:
    return (s or "").casefold().replace("ä", "a").replace("ö", "o").replace("ü", "u").replace("ß", "ss")


def _hours(h):
    return f"{h:02d}:00–{(h + 1) % 24:02d}:00" if h is not None else ""


def orders_tab(title: str, rows: list, color: str) -> Tab:
    """Every long-wait order: city A–Z → inside a city the restaurant with the most long waits first → day → time."""
    cols = ["Date", "City", "Restaurant", "Order", "Rider", "Dispatched", "Rider at\nrestaurant", "Picked up", "Waited\n(min)", "Lost\n(min)",
            "Delivered", "PTOD\n(min)", "Late?", "Note", "MotionTools"]
    t = Tab(title, len(cols), color, freeze=(1, 3))
    t.add([None] * len(cols)); t.header(0, cols)
    per = Counter((r["city"], r["rest"]) for r in rows)
    rows = sorted(rows, key=lambda r: (_ck(r["city"]), -per[(r["city"], r["rest"])], _ck(r["rest"]), r["date"], r["arr"] or dtime(0)))
    blocks, start, prev = [], 1, None
    for r in rows:
        i = t.add([r["date"], r["city"], r["rest"], r["ref"], r["rider"], r["disp"], r["arr"], r["pick"], r["wait"], max(0, r["wait"] - 5),
                   r["deliv"], r["ptod"], "late" if r["ptod"] is not None and r["ptod"] > 30 else "", r["note"],
                   Link("open ↗", r["link"]) if r.get("link") else ""])
        if r["city"] != prev and prev is not None:
            blocks.append((start, i)); start = i
        prev = r["city"]
    if rows:
        blocks.append((start, len(t.rows)))
    n = len(t.rows)
    t.style(1, n, 0, len(cols), font="Arial", size=10, valign="MIDDLE", align="CENTER")
    for c in (1, 2, 4, 13):
        t.style(1, n, c, c + 1, align="LEFT")
    t.style(1, n, 1, 2, bold=True); t.style(1, n, 8, 9, bold=True)
    t.style(1, n, 0, 1, num=DATE_F)
    for c in (5, 6, 7, 10):
        t.style(1, n, c, c + 1, num=TIME_F)
    t.style(1, n, 14, 15, color=LINK, underline=True)
    for k, (a, b) in enumerate(blocks):
        t.style(a, b, 0, len(cols), bg=BAND[k % 2])
        if k:
            t.borders.append((a, a + 1, 0, len(cols), "top"))
    t.wait_colors(1, n, 8)
    t.cond.append(((1, n, 12, 13), "TEXT_EQ", ["late"], {"textFormat": {"bold": True, "foregroundColor": _rgb(RED)}}))
    t.cond.append(((1, n, 11, 12), "NUMBER_GREATER", ["30"], {"textFormat": {"bold": True, "foregroundColor": _rgb(RED)}}))
    t.widths = dict(enumerate([84, 140, 250, 84, 120, 84, 84, 84, 70, 60, 84, 70, 60, 230, 96]))
    t.filter = (0, n, 0, len(cols))
    return t


def summary_tab(title: str, heading: str, period_text: str, refresh_text: str, rows: list, delivered: dict, weekly: bool, color: str) -> Tab:
    t = Tab(title, 10 if weekly else 9, color, freeze=(3, 0), grid=False)
    t.add([heading]); t.add([period_text, None, None, refresh_text]); t.add([])
    t.style(0, 1, 0, 1, size=16, bold=True, color=INK)
    t.style(1, 2, 0, 1, size=11, color="4B4560")
    t.style(1, 2, 3, 4, size=9, italic=True, color="6B6480")
    waits = [r["wait"] for r in rows]
    tot_del = sum(delivered.values())
    lost_h = lambda ws: sum(max(0, w - 5) for w in ws) / 60
    t.add(["Orders delivered", "Waited 10+ min", "Share", "Rider-hours lost", "Avg wait (min)", "Longest (min)"])
    t.add([tot_del, len(waits), (len(waits) / tot_del) if tot_del else 0, round(lost_h(waits), 1),
           round(sum(waits) / len(waits), 1) if waits else 0, max(waits) if waits else 0])
    t.style(3, 5, 0, 6, bg="F4F1FB", font="Arial")
    t.style(3, 4, 0, 6, size=9, color="6B6480", valign="BOTTOM", wrap=True, align="CENTER")
    t.style(4, 5, 0, 6, size=15, bold=True, color=INK, align="CENTER")
    t.style(3, 5, 0, 1, align="LEFT")
    t.style(4, 5, 1, 2, color=RED); t.style(4, 5, 3, 4, color=RED)
    for c, f in ((0, INT_F), (1, INT_F), (2, PCT_F), (3, INT_F), (4, ONE_F), (5, INT_F)):
        t.style(4, 5, c, c + 1, num=f)
    t.heights[3] = 34; t.heights[4] = 36
    t.add([])

    # by city, A–Z
    t.add(["By city (A–Z)"]); t.style(len(t.rows) - 1, len(t.rows), 0, 1, size=12, bold=True, color=INK)
    by_city = defaultdict(list)
    for r in rows:
        by_city[r["city"]].append(r)
    cities = sorted(set(delivered) | set(by_city), key=_ck)
    if weekly:
        hdr = ["City", "Worst restaurant", "Delivered", "Waited\n10+ min", "Share", "Avg wait\n(min)", "Longest\n(min)", "Rider-hours\nlost", "Its long\nwaits"]
    else:
        hdr = ["City", "Delivered", "Waited\n10+ min", "Share", "Avg wait\n(min)", "Longest\n(min)", "Rider-hours\nlost", "Worst restaurant", "Its long\nwaits"]
    h0 = t.add([]); t.header(h0, hdr)
    first = len(t.rows)
    for c in cities:
        rs = by_city[c]
        ws = [r["wait"] for r in rs]
        worst = Counter(r["rest"] for r in rs).most_common(1)
        dl = delivered.get(c, 0)
        nums = [dl, len(ws), (len(ws) / dl) if dl else 0, round(sum(ws) / len(ws), 1) if ws else None, max(ws) if ws else None, round(lost_h(ws), 1)]
        wname, wn = (worst[0][0], worst[0][1]) if worst else ("", None)
        t.add([c, wname] + nums + [wn] if weekly else [c] + nums + [wname, wn])
    last = len(t.rows)
    tot = [tot_del, len(waits), (len(waits) / tot_del) if tot_del else 0, round(sum(waits) / len(waits), 1) if waits else None,
           max(waits) if waits else None, round(lost_h(waits), 1)]
    t.add(["All cities", ""] + tot + [""] if weekly else ["All cities"] + tot + ["", ""])
    off = 2 if weekly else 1                                   # first number column
    t.style(first, last + 1, 0, 9, font="Arial", size=10, align="CENTER", valign="MIDDLE")
    t.style(first, last, 0, 1, bold=True, align="LEFT")
    t.style(first, last, off + 1, off + 2, bold=True)
    t.style(first, last + 1, 1 if weekly else 7, 2 if weekly else 8, align="LEFT")
    for k in range(first, last):
        t.style(k, k + 1, 0, 9, bg=BAND[(k - first) % 2])
    for c, f in ((off, INT_F), (off + 1, INT_F), (off + 2, PCT_F), (off + 3, ONE_F), (off + 4, INT_F), (off + 5, ONE_F)):
        t.style(first, last + 1, c, c + 1, num=f)
    t.style(last, last + 1, 0, 9, bg=INK, bold=True, color="FFFFFF")
    t.style(last, last + 1, 0, 1, align="LEFT")
    t.cond.append(((first, last, off + 1, off + 2), "GRADIENT", ["0", "max"], ("FFFFFF", "C9B8F7")))
    t.cond.append(((first, last, off + 2, off + 3), "GRADIENT3", ["0.1", "0.25", "0.4"], ("FFFFFF", "FFE7B3", "F4B6B6")))
    t.wait_colors(first, last, off + 3)
    t.add([]); t.add([])

    if not weekly:                     # yesterday: the 10 restaurants with the most long waits
        t.add(["Top 10 restaurants (most long waits)"]); t.style(len(t.rows) - 1, len(t.rows), 0, 1, size=12, bold=True, color=INK)
        h1 = t.add([]); t.header(h1, ["Restaurant", "City", "Long\nwaits", "Avg wait\n(min)", "Longest\n(min)", "Rider-h\nlost", "Worst hour"])
        f1 = len(t.rows)
        by_r = defaultdict(list)
        for r in rows:
            by_r[(r["rest"], r["city"])].append(r)
        for (rest, c), rs in sorted(by_r.items(), key=lambda x: (-len(x[1]), _ck(x[0][0])))[:10]:
            ws = [r["wait"] for r in rs]
            hrs = Counter(r["arr"].hour for r in rs if r["arr"])
            t.add([rest, c, len(ws), round(sum(ws) / len(ws), 1), max(ws), round(lost_h(ws), 1), _hours(hrs.most_common(1)[0][0] if hrs else None)])
        l1 = len(t.rows)
        t.style(f1, l1, 0, 7, font="Arial", size=10, align="CENTER", valign="MIDDLE")
        t.style(f1, l1, 0, 2, align="LEFT"); t.style(f1, l1, 0, 1, bold=True); t.style(f1, l1, 2, 3, bold=True)
        for k in range(f1, l1):
            t.style(k, k + 1, 0, 7, bg=BAND[(k - f1) % 2])
        t.style(f1, l1, 3, 4, num=ONE_F); t.style(f1, l1, 5, 6, num=ONE_F)
        t.wait_colors(f1, l1, 3)
        t.widths = dict(enumerate([200, 96, 96, 80, 88, 80, 104, 270, 88]))
    else:                              # last week: the call list per city
        t.add(["Restaurants to call — per city (3+ long waits in the week)"]); t.style(len(t.rows) - 1, len(t.rows), 0, 1, size=12, bold=True, color=INK)
        h1 = t.add([]); t.header(h1, ["City", "Restaurant", "Long\nwaits", "Days with\nlong waits", "Avg wait\n(min)", "Longest\n(min)", "Rider-hours\nlost",
                                      "Worst hour", "Worst day", "What to do"])
        f1 = len(t.rows)
        agg = defaultdict(list)
        for r in rows:
            agg[(r["city"], r["rest"])].append(r)
        keep = sorted(((k, rs) for k, rs in agg.items() if len(rs) >= 3), key=lambda x: (_ck(x[0][0]), -lost_h([r["wait"] for r in x[1]])))
        starts, prev = [], None
        ndays = max(1, len({r["date"] for r in rows}))
        for (c, rest), rs in keep:
            ws = [r["wait"] for r in rs]
            hrs = Counter(r["arr"].hour for r in rs if r["arr"])
            wh = hrs.most_common(1)[0][0] if hrs else None
            days = Counter(r["date"] for r in rs)
            per_day, avgw = len(ws) / 7, sum(ws) / len(ws)
            todo = (f"Call this week · send riders later around {wh:02d}:00" if per_day >= 3 and wh is not None else
                    f"Call · watch {wh:02d}:00" if per_day >= 1 and wh is not None else "Keep an eye on it")
            if avgw >= 15:
                todo += " · add to the Lieferando evidence pack"
            i = t.add([c, rest, len(ws), len(days), round(avgw, 1), max(ws), round(lost_h(ws), 1), _hours(wh), days.most_common(1)[0][0], todo])
            if c != prev:
                starts.append(i)
            prev = c
        l1 = len(t.rows)
        t.style(f1, l1, 0, 10, font="Arial", size=10, align="CENTER", valign="MIDDLE")
        t.style(f1, l1, 0, 2, align="LEFT"); t.style(f1, l1, 9, 10, align="LEFT")
        t.style(f1, l1, 0, 1, bold=True); t.style(f1, l1, 2, 3, bold=True)
        t.style(f1, l1, 4, 5, num=ONE_F); t.style(f1, l1, 6, 7, num=ONE_F); t.style(f1, l1, 8, 9, num=DATE_F)
        for k, a in enumerate(starts):
            b = starts[k + 1] if k + 1 < len(starts) else l1
            t.style(a, b, 0, 10, bg=BAND[k % 2])
            if k:
                t.borders.append((a, a + 1, 0, 10, "top"))
        t.wait_colors(f1, l1, 4)
        t.cond.append(((f1, l1, 6, 7), "GRADIENT", ["0", "max"], ("FFFFFF", "F2B8B5")))
        if l1 > f1:
            t.filter = (h1, l1, 0, 10)
        t.widths = dict(enumerate([180, 250, 80, 90, 80, 76, 96, 104, 90, 560]))
        _ = ndays
    t.style(0, len(t.rows), 0, t.ncols, font="Arial")
    return t


def help_tab() -> Tab:
    t = Tab("How it works", 2, "D1D5DB", grid=False)
    lines = [
        ("How this sheet works", None), ("", None),
        ("What counts", "An order where the rider waited 10 minutes or more for the food: the first rider's arrival at the restaurant → pickup. "
                        "A rider who gave up and handed the order back still counts. Orders the restaurant cancelled while the rider waited are included (see Note)."),
        ("Lost (min)", "Minutes beyond a normal 5-minute hand-over: Waited − 5. All of them together ÷ 60 = rider-hours lost."),
        ("Late?", "PTOD (dispatch → delivered) over 30 minutes."),
        ("Colours", "Waited 10–14 min = yellow · 15–19 min = orange · 20+ min = red."),
        ("Sorting", "Sorted by city A–Z. Inside a city: the restaurant with the most long waits first, then by time."),
        ("Every morning", "Quickzi Ops makes this file for yesterday and posts the download link with a short summary in the team's Slack channel. "
                          "Mondays the file also has last week (Monday–Sunday)."),
        ("Google Sheets", "Drive → New → File upload, then open the file — or in any Google Sheet: File → Import → Upload."),
        ("Privacy", "No customer data: no names, addresses, phone numbers or zip codes. Order numbers, restaurants, riders and times only. "
                    "Keep the file inside the team."),
        ("Filters", "Use Data → Create a filter view to see one city or one restaurant without changing the sheet for everybody else."),
    ]
    for a, b in lines:
        t.add([a, b])
    t.style(0, len(t.rows), 0, 2, font="Arial", size=10, valign="TOP")
    t.style(0, 1, 0, 1, size=16, bold=True, color=INK)
    t.style(2, len(t.rows), 0, 1, bold=True, color=INK)
    t.style(2, len(t.rows), 1, 2, wrap=True)
    t.widths = {0: 170, 1: 760}
    return t


def build(day: dict | None, week: dict | None) -> list:
    """day / week: {"rows": [...], "delivered": {city: n}, "label": "...", "title": "Yesterday" | "Day" | "Last week" | "Week"}."""
    tabs = []
    if day is not None:
        tabs.append(summary_tab(day["title"], "Long restaurant waits — riders who waited 10+ minutes for the food", day["label"],
                                "Made by Quickzi Ops · all cities · no customer data", day["rows"], day["delivered"], False, VIOLET))
    if week is not None:
        tabs.append(summary_tab(week["title"], f"{week['title']} — riders who waited 10+ minutes for the food", week["label"],
                                "Made by Quickzi Ops · all cities · no customer data", week["rows"], week["delivered"], True, "B9A4F5"))
    if day is not None:
        tabs.append(orders_tab(day["title"] + " orders", day["rows"], "F59E0B"))
    if week is not None:
        tabs.append(orders_tab(week["title"] + " orders", week["rows"], "FCD34D"))
    tabs.append(help_tab())
    return tabs


# ------------------------------------------------------------------ the Excel file
def to_xlsx(tabs: list) -> bytes:
    """The tabs as an .xlsx (Arial, colours, frozen headers, filters, conditional colours, links)."""
    from openpyxl import Workbook
    from openpyxl.formatting.rule import CellIsRule, ColorScaleRule
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter as L

    hx = lambda c: "%02X%02X%02X" % tuple(round(c[k] * 255) for k in ("red", "green", "blue"))
    cache = {}

    def obj(kind, key, make):
        k = (kind, key)
        if k not in cache:
            cache[k] = make()
        return cache[k]
    wb = Workbook()
    wb.remove(wb.active)
    for t in tabs:
        ws = wb.create_sheet(t.title)
        nr, nc = len(t.rows), t.ncols
        # every cell's format: the style ranges applied in order, later ones win (as in Google Sheets)
        st = [[{} for _ in range(nc)] for _ in range(nr)]
        for (r0, r1, c0, c1, f, _fields) in t.styles:
            flat = {}
            if "backgroundColor" in f:
                flat["bg"] = hx(f["backgroundColor"])
            tf = f.get("textFormat", {})
            for k in ("bold", "italic", "fontSize", "underline", "fontFamily"):
                if k in tf:
                    flat[k] = tf[k]
            if "foregroundColor" in tf:
                flat["color"] = hx(tf["foregroundColor"])
            for k in ("horizontalAlignment", "verticalAlignment", "wrapStrategy"):
                if k in f:
                    flat[k] = f[k]
            if "numberFormat" in f:
                flat["num"] = f["numberFormat"]["pattern"]
            for r in range(r0, min(r1, nr)):
                row = st[r]
                for c in range(c0, min(c1, nc)):
                    row[c].update(flat)
        for r, vals in enumerate(t.rows):
            for c, v in enumerate(vals):
                s_ = st[r][c]
                if v is None and not s_:
                    continue
                link = None
                if isinstance(v, Link):
                    v, link = v.text, v.url
                cell = ws.cell(row=r + 1, column=c + 1, value=v if v != "" else None)
                if link:
                    cell.hyperlink = link
                if not s_:
                    continue
                fk = (s_.get("fontFamily", "Arial"), s_.get("fontSize", 10), bool(s_.get("bold")), bool(s_.get("italic")),
                      bool(s_.get("underline")), s_.get("color"))
                cell.font = obj("font", fk, lambda: Font(name=fk[0], size=fk[1], bold=fk[2], italic=fk[3],
                                                         underline="single" if fk[4] else None, color=fk[5]))
                if s_.get("bg"):
                    cell.fill = obj("fill", s_["bg"], lambda: PatternFill("solid", fgColor=s_["bg"]))
                ak = (s_.get("horizontalAlignment"), s_.get("verticalAlignment"), s_.get("wrapStrategy") == "WRAP")
                if any(ak):
                    cell.alignment = obj("al", ak, lambda: Alignment(horizontal=(ak[0] or "general").lower(),
                                                                     vertical={"MIDDLE": "center"}.get(ak[1], (ak[1] or "bottom").lower()),
                                                                     wrap_text=ak[2]))
                if s_.get("num"):
                    cell.number_format = s_["num"]
        side = Side(style="medium", color=VIOLET)
        for (r0, _r1, c0, c1, _side) in t.borders:
            for c in range(c0, c1):
                ws.cell(row=r0 + 1, column=c + 1).border = Border(top=side)
        for (r0, r1, c0, c1), kind, vals, f in t.cond:
            if r1 <= r0:
                continue
            ref = f"{L(c0 + 1)}{r0 + 1}:{L(c1)}{r1}"
            if kind == "GRADIENT":
                ws.conditional_formatting.add(ref, ColorScaleRule(start_type="num", start_value=float(vals[0]), start_color=f[0],
                                                                  end_type="max", end_color=f[1]))
            elif kind == "GRADIENT3":
                ws.conditional_formatting.add(ref, ColorScaleRule(start_type="num", start_value=float(vals[0]), start_color=f[0],
                                                                  mid_type="num", mid_value=float(vals[1]), mid_color=f[1],
                                                                  end_type="num", end_value=float(vals[2]), end_color=f[2]))
            else:
                kw = {}
                if "backgroundColor" in f:
                    kw["fill"] = PatternFill("solid", fgColor=hx(f["backgroundColor"]))
                if "textFormat" in f:
                    tf = f["textFormat"]
                    kw["font"] = Font(name="Arial", bold=tf.get("bold"), color=hx(tf["foregroundColor"]) if "foregroundColor" in tf else None)
                op = {"NUMBER_GREATER_THAN_EQ": "greaterThanOrEqual", "NUMBER_BETWEEN": "between", "NUMBER_GREATER": "greaterThan", "TEXT_EQ": "equal"}[kind]
                ws.conditional_formatting.add(ref, CellIsRule(operator=op, formula=[f'"{vals[0]}"'] if kind == "TEXT_EQ" else list(vals), **kw))
        for c, px in t.widths.items():
            ws.column_dimensions[L(c + 1)].width = round(px / 7, 1)
        for r, px in t.heights.items():
            ws.row_dimensions[r + 1].height = round(px * 0.75, 1)
        fr, fc = t.freeze
        if fr or fc:
            ws.freeze_panes = f"{L(fc + 1)}{fr + 1}"
        if t.filter and t.filter[1] > t.filter[0] + 1:
            r0, r1, c0, c1 = t.filter
            ws.auto_filter.ref = f"{L(c0 + 1)}{r0 + 1}:{L(c1)}{r1}"
        ws.sheet_view.showGridLines = t.grid
        ws.sheet_properties.tabColor = t.color
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
